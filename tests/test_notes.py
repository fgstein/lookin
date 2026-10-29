"""A workspace's notes: the Markdown they are kept as, the recap Lookin
writes into them, the tab that edits them, the MCP tools that read and
write them, and the reports that carry them.

What is pinned: the editor's HTML and the stored Markdown go round without
losing or inventing anything; nothing unsafe survives a save; the recap is
the only part a regeneration touches; a save from an editor that has not
seen the latest text never overwrites it; and the notes come first in
every report format.
"""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest

from lookin.core import agent_actions as actions
from lookin.core import agent_outputs as outputs
from lookin.core import agent_tools as tools
from lookin.core import library as lib
from lookin.core import links
from lookin.core import markdown_lite as md
from lookin.core import notes as nt

EXAMPLE = Path(__file__).resolve().parent.parent / "src" / "lookin" / "data" / "example"

NOTE = """# My workspace

<!-- lookin:recap -->

## Performance summary

- **baseline_cnn** -- accuracy 0.639
  - nested *point*
- run_name\\_v2: `code` and [a link](https://mlflow.lan/#/runs/1)

<!-- /lookin:recap -->

## Notes

We ship **resnet**, not *baseline*. ~~Maybe~~ surely.
A second line.

1. relabel the cats
2. check the dogs

> a quote

---
"""


# ---- Markdown, both ways ----------------------------------------------------

def test_markdown_survives_the_editor_round_trip():
    html = md.to_html(NOTE)
    assert md.to_html(md.from_html(html)) == html
    assert md.from_html(md.to_html(md.from_html(html))) == md.from_html(html)
    assert '<div data-region="recap">' in html and "<em>point</em>" in html and "<s>Maybe</s>" in html


def test_what_a_browser_editor_produces_is_read_as_markdown():
    html = ('<div>hello <b>bold</b> and <span style="font-weight: 700">heavy</span></div><div><br></div>'
            '<div>x<i>it</i>&nbsp;y</div><ul><li>a<ul><li>b</li></ul></li></ul><p><ul><li>in p</li></ul></p>')
    assert md.from_html(html) == "hello **bold** and **heavy**\n\nx*it* y\n\n- a\n  - b\n\n- in p\n"


def test_nothing_unsafe_survives_a_save():
    saved = md.from_html('<p><a href="javascript:alert(1)">click</a><script>steal()</script>'
                         '<img src=x onerror=alert(1)></p>')
    assert saved.strip() == "click"
    shown = md.to_html("[x](javascript:alert(1)) <script>alert(1)</script> *ok*")
    assert "<script>" not in shown and "javascript:" not in shown and "<em>ok</em>" in shown


def test_text_that_looks_like_markup_stays_text():
    typed = md.from_html("<p>2 * 3 * 4, # not a title, snake_case_name, [not a link]</p><p>- not a list</p>")
    assert md.to_html(typed) == ("<p>2 * 3 * 4, # not a title, snake_case_name, [not a link]</p>\n"
                                 "<p>- not a list</p>")


def test_a_report_demotes_the_notes_headings_and_drops_the_markers():
    out = md.for_report(NOTE)
    assert out.startswith("### My workspace") and "#### Notes" in out and "lookin:recap" not in out
    assert "relabel the cats" in md.to_plain(NOTE) and "**" not in md.to_plain(NOTE)


# ---- the recap ----------------------------------------------------------------

def test_regenerating_touches_the_recap_only():
    again = nt.with_recap(NOTE, "## Performance summary\n\n- new numbers")
    assert "new numbers" in again and "baseline_cnn" not in again
    assert again.split(md.region_close("recap"))[1] == NOTE.split(md.region_close("recap"))[1]
    # Markers deleted by the user: a fresh recap under the title, the rest kept.
    bare = "# Title\n\nmy text\n"
    assert nt.with_recap(bare, "R").startswith("# Title\n\n<!-- lookin:recap -->\n\nR\n\n<!-- /lookin:recap -->")
    assert nt.with_recap("no title", "R").endswith("no title")


# ---- over MCP -----------------------------------------------------------------

@pytest.fixture
def built(tmp_path, monkeypatch):
    lib.set_library_root(tmp_path / "library")
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path / "config"))
    monkeypatch.setenv("XDG_CACHE_HOME", str(tmp_path / "cache"))
    links.set_base_url("http://lookin.lan:8560")
    tools.clear_cache()
    example = tmp_path / "ex"
    shutil.copytree(EXAMPLE, example)
    ws = actions.create_workspace("Example", dataset_path=str(example / "dataset.csv"))
    actions.add_experiment(ws["id"], str(example / "baseline_cnn_results.csv"))
    actions.add_experiment(ws["id"], str(example / "finetuned_resnet_results.csv"), name="resnet")
    yield ws["id"]
    lib.set_library_root(None)
    links.set_base_url("")
    tools.clear_cache()


def test_an_agent_reads_and_writes_the_notes(built):
    first = tools.get_notes(built)
    assert first["exists"] is False and first["warnings"] and first["open_url"].endswith("tab=notes")
    started = actions.regenerate_notes_recap(built)
    assert started["started"] and "resnet" in started["recap"]
    text = tools.get_notes(built)["notes"]
    assert text.startswith("# Example") and md.region_open("recap") in text
    actions.write_notes(built, "## Agent findings\n\n- cats are hard")
    text = tools.get_notes(built)["notes"]
    assert text.index("Performance summary") < text.index("Agent findings")
    with pytest.raises(ValueError, match="Refusing"):
        actions.write_notes(built, "  ", mode="replace")
    again = actions.regenerate_notes_recap(built)
    assert not again["started"] and "Agent findings" in tools.get_notes(built)["notes"]


def test_the_notes_open_every_report_format(built):
    actions.write_notes(built, "# Example\n\nWe ship **resnet**.", mode="replace")
    for fmt in ("html", "md", "pptx"):
        out = outputs.generate_report(built, format=fmt, snapshot=False)
        path = Path(out["path"])
        assert path.suffix == f".{fmt}"
        if fmt == "pptx":
            from pptx import Presentation
            text = " ".join(shape.text_frame.text for slide in Presentation(str(path)).slides
                            for shape in slide.shapes if shape.has_text_frame)
            assert "We ship resnet." in text
            continue
        body = path.read_text(encoding="utf-8")
        assert body.index("Notes") < body.index("Summary") < body.index("Per-class charts")
        assert ("<strong>resnet</strong>" if fmt == "html" else "We ship **resnet**.") in body
    without = Path(outputs.generate_report(built, format="md", snapshot=False, notes=False)["path"])
    assert "We ship" not in without.read_text(encoding="utf-8")


def test_the_notes_travel_with_an_exported_workspace(built, tmp_path):
    actions.write_notes(built, "kept", mode="replace")
    workspace = next(w for w in lib.list_workspaces() if w.id == built)
    ok, _ = lib.export_workspace(workspace, tmp_path / "ws.zip")
    assert ok and "notes.md" in zipfile.ZipFile(tmp_path / "ws.zip").namelist()


# ---- the tab ------------------------------------------------------------------

def test_the_tab_keeps_a_note_and_never_saves_over_a_newer_one(built, tmp_path):
    from lookin.web.render import boot
    from lookin.web.session import RequestScope
    from lookin.web.server import WebApp

    session = WebApp().store.get("n" * 24, create=True)
    with RequestScope(session):
        boot(session, {})
    workspace = next(w for w in lib.list_workspaces() if w.id == built)
    view = session.view("notes")
    # A note written before the session became a workspace goes into it.
    view.text, view.started, view.loaded_for = "# Draft\n\nmine\n", True, ""
    session.current_workspace = workspace
    view._sync()
    assert nt.read(workspace) == "# Draft\n\nmine\n"

    view.on_save(html=md.to_html("# Draft\n\nedited"), rev=view.rev)
    assert "edited" in nt.read(workspace)
    # An editor that has not shown the latest text yet: its save is dropped.
    view.on_save(html="<p>stale</p>", rev=view.rev - 1)
    assert "stale" not in nt.read(workspace)
    # Written elsewhere (an agent): shown, not overwritten by the next save.
    actions.write_notes(built, "from the agent")
    rev = view.rev
    view.on_save(html="<p>typed before seeing it</p>", rev=rev)
    assert "from the agent" in nt.read(workspace) and "typed before" not in nt.read(workspace)
    assert view.rev == rev + 1 and "from the agent" in view.text
