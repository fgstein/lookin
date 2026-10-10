"""The PowerPoint rendering of the report.

`PptxReport` overrides `Report`'s content methods to build slides instead
of HTML, so report_builder.py's section-building logic (tested against
HTML in test_report.py) is exercised again here against a real .pptx --
same sections, different writer.
"""
from __future__ import annotations

from pptx import Presentation

from lookin.core.report_builder import build_report, default_report_name
from lookin.core.report_pptx import PptxReport


def _all_text(prs: Presentation) -> str:
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
    return "\n".join(chunks)


def test_a_report_is_a_real_presentation_with_slides(tmp_path, loaded_state):
    report = build_report(loaded_state, ["exp"], report_cls=PptxReport)
    dest = report.write(tmp_path / "out.pptx")
    assert dest.exists()

    prs = Presentation(dest)
    assert len(prs.slides) > 1
    assert not list(tmp_path.glob(".*tmp*")), "the temp file is replaced, not left behind"


def test_the_report_covers_the_loaded_experiment(tmp_path, loaded_state):
    report = build_report(loaded_state, ["exp"], report_cls=PptxReport)
    prs = Presentation(report.write(tmp_path / "out.pptx"))
    text = _all_text(prs)
    assert "Summary" in text and "Per class" in text and "Confusion matrices" in text
    assert "exp" in text
    assert "Accuracy" in text


def test_provenance_says_whether_reviewed_labels_were_applied(tmp_path, loaded_state):
    from lookin.core.review import STATUS_EXCLUDE

    report = build_report(loaded_state, ["exp"], report_cls=PptxReport)
    text = _all_text(Presentation(report.write(tmp_path / "a.pptx")))
    assert "Reviewed labels applied" in text

    loaded_state.review.create_version("v1")
    loaded_state.review.decide("a", STATUS_EXCLUDE)
    loaded_state.use_reviewed_labels = True
    report = build_report(loaded_state, ["exp"], report_cls=PptxReport)
    text = _all_text(Presentation(report.write(tmp_path / "b.pptx")))
    assert "cleaned pool" in text
    assert "v1" in text


def test_a_confusion_matrix_figure_is_embedded_as_a_picture(tmp_path, loaded_state):
    report = build_report(loaded_state, ["exp"], report_cls=PptxReport)
    prs = Presentation(report.write(tmp_path / "out.pptx"))
    pictures = [shape for slide in prs.slides for shape in slide.shapes if shape.shape_type == 13]
    assert pictures, "the confusion matrix should be a picture on some slide"


def test_the_comparison_section_is_optional_and_directional(tmp_path, loaded_state, results_csv):
    loaded_state.add_experiment(str(results_csv), "exp2")

    without = build_report(loaded_state, ["exp", "exp2"], report_cls=PptxReport)
    text = _all_text(Presentation(without.write(tmp_path / "without.pptx")))
    assert "Compare:" not in text

    with_pair = build_report(loaded_state, ["exp", "exp2"], comparison_pair=("exp", "exp2"),
                             report_cls=PptxReport)
    text = _all_text(Presentation(with_pair.write(tmp_path / "with.pptx")))
    assert "Compare: exp to exp2" in text
    assert "used to be right are now wrong" in text


def test_a_long_table_is_paginated_across_slides(tmp_path, loaded_state):
    """14 rows per slide (MAX_TABLE_ROWS) -- a class list longer than that
    must not silently overflow one slide."""
    import pandas as pd

    from lookin.core.data_manager import AppState, ColumnConfig
    from lookin.core.report_pptx import MAX_TABLE_ROWS

    n_classes = MAX_TABLE_ROWS + 5
    labels = [f"class_{i}" for i in range(n_classes)]
    rows = pd.DataFrame({
        "uuid": [str(i) for i in range(n_classes)],
        "crop": [f"crops/{i}.png" for i in range(n_classes)],
        "label": labels,
    })
    dataset_path = tmp_path / "many_classes.csv"
    rows.to_csv(dataset_path, index=False)

    scores = pd.DataFrame({"uuid": [str(i) for i in range(n_classes)],
                           **{label: [1.0 if l == label else 0.0 for l in labels] for label in labels}})
    results_path = tmp_path / "many_classes_exp_results.csv"
    scores.to_csv(results_path, index=False)

    state = AppState()
    state.load_dataset(str(dataset_path), ColumnConfig())
    state.add_experiment(str(results_path), "exp")

    report = build_report(state, ["exp"], report_cls=PptxReport)
    text = _all_text(Presentation(report.write(tmp_path / "long.pptx")))
    assert "(1/2)" in text and "(2/2)" in text


def test_the_default_name_carries_the_dataset_and_the_date(loaded_state):
    name = default_report_name(loaded_state, extension="pptx")
    assert name.startswith("dataset_report_") and name.endswith(".pptx")


def test_a_report_with_no_experiment_is_refused(loaded_state):
    import pytest

    with pytest.raises(ValueError):
        build_report(loaded_state, [], report_cls=PptxReport)
